import pandas as pd
import glob
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

# 1. Поиск и объединение CSV
csv_files = glob.glob("results_*.csv")
dataframes = []
for file in csv_files:
    df = pd.read_csv(file)
    dataframes.append(df)

combined_df = pd.concat(dataframes, ignore_index=True)
combined_df = combined_df.sort_values(['model', 'prompt_id']).reset_index(drop=True)

# 2. Добавление типов задач и длины ответов
prompt_types = {
    1: "Объяснение", 2: "Классификация", 3: "Суммирование", 4: "Извлечение данных",
    5: "Генерация кода", 6: "Поиск ошибки", 7: "Логика", 
    8: "Информационная безопасность", 9: "Деловой стиль", 10: "Жесткий формат"
}
combined_df['prompt_type'] = combined_df['prompt_id'].map(prompt_types)
combined_df['answer_length_chars'] = combined_df['answer'].str.len()
combined_df['answer_length_words'] = combined_df['answer'].str.split().str.len()

# 3. Инициализация колонок для оценок (чтобы не было KeyError)
score_cols = ['score_correctness', 'score_completeness', 'score_instruction', 'score_usefulness', 'score_total']
for col in score_cols:
    if col not in combined_df.columns:
        combined_df[col] = None

# 4. Вывод сводки в консоль
print("Сводная таблица по моделям:")
summary = combined_df.groupby('model').agg({
    'latency': 'mean',
    'answer_length_chars': 'mean',
    'error': lambda x: x.notna().sum()
}).rename(columns={
    'latency': 'Средняя задержка',
    'answer_length_chars': 'Средняя длина ответа',
    'error': 'Количество ошибок'
})
print(summary)

# 5. Создание Excel
wb = Workbook()
ws_data = wb.active
ws_data.title = "Данные"

column_order = ['model', 'prompt_id', 'prompt_type', 'latency', 'answer', 'error', 
                'answer_length_chars', 'answer_length_words', 'score_correctness', 
                'score_completeness', 'score_instruction', 'score_usefulness', 'score_total']

data_for_excel = combined_df[column_order].copy()
headers = ['Модель', 'Номер задачи', 'Тип задачи', 'Задержка', 'Ответ', 'Ошибка', 
           'Длина ответа (символы)', 'Длина ответа (слова)', 'Корректность', 
           'Полнота', 'Инструкция', 'Полезность', 'Общий балл']

ws_data.append(headers)
for r in dataframe_to_rows(data_for_excel, index=False, header=False):
    ws_data.append(r)

# Форматирование
for cell in ws_data[1]:
    cell.font = Font(bold=True)
ws_data.freeze_panes = "A2"

# Формула для score_total (колонка M = 13)
for row in range(2, ws_data.max_row + 1):
    ws_data.cell(row=row, column=13).value = f"=SUM(I{row}:L{row})"

# Валидация данных для оценок (колонки I, J, K, L)
for col in [9, 10, 11, 12]:
    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    ws_data.add_data_validation(dv)
    for row in range(2, ws_data.max_row + 1):
        dv.add(ws_data.cell(row=row, column=col))

# 6. Лист "Сводка"
ws_summary = wb.create_sheet(title="Сводка")
summary_headers = ['Модель', 'Средний балл', 'Средняя корректность', 
                   'Средняя полнота', 'Средняя инструкция', 'Средняя полезность', 
                   'Средняя задержка', 'Средняя длина ответа', 'Количество ошибок']
ws_summary.append(summary_headers)
for cell in ws_summary[1]:
    cell.font = Font(bold=True)

models = combined_df['model'].unique()
for idx, model in enumerate(models, start=2):
    ws_summary.cell(row=idx, column=1).value = model
    ws_summary.cell(row=idx, column=2).value = f'=ROUND(AVERAGEIF(Данные!$A:$A, A{idx}, Данные!$M:$M), 2)'
    ws_summary.cell(row=idx, column=3).value = f'=ROUND(AVERAGEIF(Данные!$A:$A, A{idx}, Данные!$I:$I), 2)'
    ws_summary.cell(row=idx, column=4).value = f'=ROUND(AVERAGEIF(Данные!$A:$A, A{idx}, Данные!$J:$J), 2)'
    ws_summary.cell(row=idx, column=5).value = f'=ROUND(AVERAGEIF(Данные!$A:$A, A{idx}, Данные!$K:$K), 2)'
    ws_summary.cell(row=idx, column=6).value = f'=ROUND(AVERAGEIF(Данные!$A:$A, A{idx}, Данные!$L:$L), 2)'
    ws_summary.cell(row=idx, column=7).value = f'=ROUND(AVERAGEIF(Данные!$A:$A, A{idx}, Данные!$D:$D), 3)'
    ws_summary.cell(row=idx, column=8).value = f'=ROUND(AVERAGEIF(Данные!$A:$A, A{idx}, Данные!$G:$G), 1)'
    ws_summary.cell(row=idx, column=9).value = f'=COUNTIFS(Данные!$A:$A, A{idx}, Данные!$F:$F, "<>")'

tab = Table(displayName="SummaryTable", ref=f"A1:I{ws_summary.max_row}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws_summary.add_table(tab)

wb.save("results_all.xlsx")

# 7. Создание Markdown для чтения
with open('results_review.md', 'w', encoding='utf-8') as f:
    f.write("# Анализ ответов моделей\n\n")
    for model in models:
        f.write(f"## Модель: {model}\n\n")
        model_data = combined_df[combined_df['model'] == model]
        for p_type in prompt_types.values():
            f.write(f"### {p_type}\n\n")
            type_data = model_data[model_data['prompt_type'] == p_type]
            if len(type_data) > 0:
                for _, row in type_data.iterrows():
                    ans = str(row['answer'])
                    ans = ans[:500] + "..." if len(ans) > 500 else ans
                    f.write(f"- **Ответ**: {ans}\n")
            else:
                f.write("- Нет данных\n")
            f.write("\n")

print("\nАнализ завершен. Созданы файлы: results_all.xlsx и results_review.md")